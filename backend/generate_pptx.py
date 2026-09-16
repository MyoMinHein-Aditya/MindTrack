from pptx import Presentation
from pptx.util import Inches, Pt

prs = Presentation()

def add_slide(title_text, content_text):
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    content = slide.placeholders[1]
    
    title.text = title_text
    content.text = content_text

# Slide 1
add_slide(
    "MindTrack: Mental Health & Well-being Surveillance", 
    "Problem Statement: Mental health and well-being surveillance, assessment and tracking solution among children\n\n"
    "Tagline: Early Detection. Timely Support. Safe Childhoods.\n\n"
    "Team Name: Team GammaRay\n\n"
    "Core Idea\nMindTrack is a child-centric digital platform designed to support periodic surveillance, age-appropriate assessment and longitudinal tracking of children's mental health and well-being, while helping trained school counselors identify children who may need additional support."
)

# Slide 2
add_slide(
    "PROPOSED SOLUTION",
    "CORE PROBLEM\n"
    "- Children's mental health can change gradually and go unnoticed.\n"
    "- Children may lack vocabulary to communicate emotional difficulties.\n"
    "- Need for structured survey, assess and track system.\n\n"
    "PROPOSED SOLUTION\n"
    "- Child-friendly digital check-ins.\n"
    "- Age-appropriate questions.\n"
    "- Compares current responses with historical trends.\n"
    "- Classifies cases for counselors to follow up.\n\n"
    "3 CORE FUNCTIONS\n"
    "SURVEILLANCE | ASSESSMENT | TRACKING"
)

# Slide 3
add_slide(
    "TECHNICAL APPROACH",
    "FRONTEND\n"
    "- React.js, Vite, Framer Motion, Gamified UX\n\n"
    "BACKEND\n"
    "- FastAPI, Python, REST APIs, PostgreSQL\n\n"
    "AI / ANALYTICS ENGINE\n"
    "- Groq API, Llama 3 8B, Structured response analysis\n\n"
    "3-TIER TRACKING ENGINE\n"
    "TIER 1 (STABLE) -> TIER 2 (WATCHLIST) -> TIER 3 (HIGH PRIORITY)\n\n"
    "SAFETY PRINCIPLE\n"
    "'AI assists screening. Trained professionals assess and intervene.'"
)

# Slide 4
add_slide(
    "FEASIBILITY AND VIABILITY",
    "FEASIBILITY\n"
    "- Software-first solution; no specialized hardware MVP.\n"
    "- PostgreSQL supports structured longitudinal scaling.\n\n"
    "MITIGATION STRATEGIES\n"
    "- Age-appropriate, structured questions.\n"
    "- Combine assessment with historical trends.\n"
    "- Strict RBAC and school-level data isolation.\n\n"
    "ROADMAP\n"
    "MVP -> PILOT -> DISTRICT SCALE -> STATE SCALE"
)

# Slide 5
add_slide(
    "IMPACT AND BENEFITS",
    "CHILD: Simple, non-intimidating reflection; early support.\n"
    "COUNSELOR: Prioritized list; longitudinal graphs; tracking.\n"
    "ADMINISTRATION: Anonymized aggregate trends.\n\n"
    "TARGET OUTCOMES\n"
    "- Earlier identification of concerns.\n"
    "- Better visibility into gradual changes.\n"
    "- Reduced dependence on observation alone.\n\n"
    "KEY IMPACT STATEMENT\n"
    "'Surveillance finds the signal. Assessment understands it. Tracking reveals the trend.'"
)

# Slide 6
add_slide(
    "RESEARCH, ETHICS & REFERENCES",
    "1. CHILD-CENTRIC ASSESSMENT\n"
    "- Age-appropriate, gamified, non-diagnostic.\n\n"
    "2. PRIVACY-FIRST DESIGN\n"
    "- RBAC, data isolation, auditable profiles.\n\n"
    "3. HUMAN-IN-THE-LOOP SAFETY\n"
    "- AI never replaces the counselor.\n\n"
    "REFERENCES\n"
    "- India's Digital Personal Data Protection framework.\n"
    "- Institutional child-protection procedures."
)

prs.save("MindTrack_SIH_Presentation.pptx")
print("Saved!")
